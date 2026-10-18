from io import BytesIO

from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.services.database import approved_rows

HEADERS = (
    ("filename", "שם הקובץ"),
    ("page", "עמוד"),
    ("form_type", "סוג טופס"),
    ("first_name", "שם פרטי"),
    ("last_name", "שם משפחה"),
    ("id_number", "תעודת זהות"),
    ("center_number", "מספר מרכזי"),
    ("document_date", "תאריך קבלת הדגימה"),
    ("sender", "גורם שולח"),
    ("commitment_number", "מספר התחייבות"),
)


def _safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value.lstrip(" \t\r\n")[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def build_approved_workbook(document_ids: list[str] | None = None) -> BytesIO:
    rows = approved_rows(document_ids)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "טפסים מאושרים"
    worksheet.append([label for _, label in HEADERS])
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="287EAE")
        cell.alignment = Alignment(horizontal="right")
    for row in rows:
        values = []
        for key, _ in HEADERS:
            value = row.get(key, "")
            cell_value = ", ".join(map(str, value)) if isinstance(value, list) else value
            values.append(_safe_cell(cell_value))
        worksheet.append(values)
    for column in worksheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 36)
        worksheet.column_dimensions[column[0].column_letter].width = width
    worksheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
