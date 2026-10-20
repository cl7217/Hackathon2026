import unittest
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from app.services.excel import build_approved_workbook


class ExcelExportTests(unittest.TestCase):
    @patch(
        "app.services.excel.approved_rows",
        return_value=[
            {
                "filename": "sample.pdf",
                "page": 1,
                "form_type": "clalit",
                "first_name": "דנה",
                "last_name": "כהן",
                "id_number": "123456782",
                "center_number": ["00123", "00456"],
                "document_date": "20/10/2026",
                "sender": "כללית",
                "commitment_number": ["300123", "300456"],
            }
        ],
    )
    def test_export_creates_one_repeated_data_row_per_commitment(self, _approved_rows):
        workbook_stream = build_approved_workbook()
        worksheet = load_workbook(BytesIO(workbook_stream.getvalue())).active

        self.assertEqual(worksheet.max_row, 3)
        self.assertEqual(worksheet.cell(row=2, column=4).value, "דנה")
        self.assertEqual(worksheet.cell(row=3, column=4).value, "דנה")
        self.assertEqual(worksheet.cell(row=2, column=7).value, "00123, 00456")
        self.assertEqual(worksheet.cell(row=3, column=7).value, "00123, 00456")
        self.assertEqual(worksheet.cell(row=2, column=10).value, "300123")
        self.assertEqual(worksheet.cell(row=3, column=10).value, "300456")


if __name__ == "__main__":
    unittest.main()
